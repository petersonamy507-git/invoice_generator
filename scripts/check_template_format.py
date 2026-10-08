"""Generate filled templates 6-14 and print structure for format review."""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document

from backend.app.models import InvoiceTask, WordInvoiceData
from backend.app.services.word_invoice_generator import generate_word_invoice
from backend.app.services.word_to_pdf import convert_docx_batch_to_pdf


def main() -> None:
    out = ROOT / "Data" / "word" / "output" / "_format_check"
    out.mkdir(parents=True, exist_ok=True)

    tasks = [
        InvoiceTask(main_task="Services", sub_task="Service Alpha One", amount=1000),
        InvoiceTask(main_task="Services", sub_task="Service Beta Two", amount=1500),
        InvoiceTask(main_task="Services", sub_task="Service Gamma Three", amount=2000),
    ]

    def make(n: int) -> WordInvoiceData:
        return WordInvoiceData(
            person_name="Test Person",
            company_name="Test Co",
            email="t@example.com",
            total_amount=4500,
            invoice_number=n,
            bank="HBL",
            iban="PK00TEST1234567890",
            branch_code="0123",
            document_invoice_no="HH-008",
            tasks=tasks,
        )

    items: list[tuple[bytes, str]] = []
    for n in range(6, 15):
        data = make(n)
        docx_bytes, filename = generate_word_invoice(data)
        path = out / f"tpl{n}_{filename}"
        path.write_bytes(docx_bytes)
        items.append((docx_bytes, filename))
        d = Document(str(path))
        print("\n======== TEMPLATE", n, filename)
        print("paras", len(d.paragraphs), "tables", len(d.tables))
        for i, para in enumerate(d.paragraphs):
            t = para.text.strip()
            if t:
                print(f"P{i}: {t[:140]}")
        for ti, table in enumerate(d.tables):
            cols = len(table.columns) if table.rows else 0
            print(f"-- table {ti} {len(table.rows)}x{cols}")
            for ri, row in enumerate(table.rows):
                cells = [
                    c.text.strip().replace("\n", " | ")[:50] for c in row.cells
                ]
                print(f"  R{ri}: {cells}")

    for pdf_bytes, pdf_name in convert_docx_batch_to_pdf(items):
        (out / pdf_name).write_bytes(pdf_bytes)
        print("PDF", pdf_name, len(pdf_bytes))
    print("done", out)


if __name__ == "__main__":
    main()

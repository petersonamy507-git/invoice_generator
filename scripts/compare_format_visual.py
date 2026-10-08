"""Generate filled invoices and render original vs generated first page as PNG."""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import fitz
from docx import Document

from backend.app.models import InvoiceTask, WordInvoiceData
from backend.app.services.word_invoice_generator import generate_word_invoice
from backend.app.services.word_templates import template_path
from backend.app.services.word_to_pdf import convert_docx_batch_to_pdf

OUT = ROOT / "Data" / "word" / "output" / "_format_cmp"
OUT.mkdir(parents=True, exist_ok=True)

SAMPLES = {
    6: Path(r"d:\AI_Cache\temp\1790191007J-RavoteckInc.docx"),
    7: Path(r"d:\AI_Cache\temp\1790191007Z-IgnitaiInc.docx"),
    8: Path(r"d:\AI_Cache\temp\1790191009M-CoretechifyInc.docx"),
    9: Path(r"d:\AI_Cache\temp\1790191009S-EcomifyInc.docx"),
    10: Path(r"d:\AI_Cache\temp\1790191009Y-CozyHomeEssentials.docx"),
    11: Path(r"d:\AI_Cache\temp\1790191010S-BeecodifyInc.docx"),
    12: Path(r"d:\AI_Cache\temp\1790191011E-BravixTechnologiesInc.docx"),
    13: Path(r"d:\AI_Cache\temp\1790191011Z-AlphaDigitalInc (1).docx"),
    14: Path(r"d:\AI_Cache\temp\1790191007G-SynergoInc.docx"),
}


def render_pdf(pdf_bytes: bytes, png_path: Path, zoom: float = 1.5) -> None:
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    page = doc[0]
    pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
    png_path.write_bytes(pix.tobytes("png"))
    doc.close()


def main() -> None:
    tasks = [
        InvoiceTask("S", "Property Valuation & Pricing Strategy", 1200),
        InvoiceTask("S", "Professional Photography", 1500),
        InvoiceTask("S", "Online Advertising + Social Media Ads", 1800),
    ]

    # Convert originals once
    orig_items: list[tuple[bytes, str]] = []
    for n, path in SAMPLES.items():
        tpl = template_path(n)
        print(n, "tpl", tpl.name, "sample_same", path.read_bytes() == tpl.read_bytes())
        orig_items.append((path.read_bytes(), f"orig_{n}.docx"))

    for pdf_bytes, name in convert_docx_batch_to_pdf(orig_items):
        (OUT / name).write_bytes(pdf_bytes)
        render_pdf(pdf_bytes, OUT / name.replace(".pdf", ".png"))
        print("orig pdf", name, len(pdf_bytes))

    # Generated
    gen_items: list[tuple[bytes, str]] = []
    for n in SAMPLES:
        data = WordInvoiceData(
            person_name="Daniel Gallego",
            company_name="Client Co",
            email="a@b.com",
            total_amount=4500,
            invoice_number=n,
            bank="HBL Bank",
            iban="0123 4567 8901",
            branch_code="0011",
            document_invoice_no="000002",
            tasks=tasks,
        )
        docx_bytes, _ = generate_word_invoice(data)
        (OUT / f"gen_{n}.docx").write_bytes(docx_bytes)
        gen_items.append((docx_bytes, f"gen_{n}.docx"))
        d = Document(OUT / f"gen_{n}.docx")
        print("gen", n, "tables", len(d.tables), "paras", len(d.paragraphs))

    for pdf_bytes, name in convert_docx_batch_to_pdf(gen_items):
        (OUT / name).write_bytes(pdf_bytes)
        render_pdf(pdf_bytes, OUT / name.replace(".pdf", ".png"))
        print("gen pdf", name, len(pdf_bytes))

    print("done", OUT)


if __name__ == "__main__":
    main()

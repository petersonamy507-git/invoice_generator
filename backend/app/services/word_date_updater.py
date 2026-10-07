import re
import shutil
from datetime import date
from pathlib import Path

from docx import Document

from backend.app.services.validation import ValidationError

INLINE_DATE_LABELS = re.compile(
    r"(?i)(Dated\s*:?\s*|Invoice Date\s*:?\s*|Date\s*:?\s*|Pay by\s*:?\s*)(.+)$"
)
DATE_VALUE = re.compile(
    r"(\d{1,2}\s*/\s*\d{1,2}\s*/\s*\d{4}|"
    r"(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|"
    r"Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)"
    r"[\s,]+\d{1,2}[\s,]+\d{4}|"
    r"\d{1,2}\s+(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|"
    r"Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)"
    r"(?:\s+\d{4})?)"
)


def format_invoice_date(d: date | None = None) -> str:
    d = d or date.today()
    return f"{d.strftime('%b')} {d.day}, {d.year}"


def _is_layout_only_text(text: str) -> bool:
    return not text.strip() or bool(re.fullmatch(r"[\s\t]+", text))


def _split_outer_whitespace(text: str) -> tuple[str, str, str]:
    core = text.strip()
    if not core:
        return "", text, ""
    start = text.index(core)
    end = start + len(core)
    return text[:start], core, text[end:]


def _replace_paragraph_text(paragraph, new_text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)


def _replace_date_in_text(text: str, new_date: str) -> tuple[str, bool]:
    match = DATE_VALUE.search(text)
    if match:
        return text[: match.start()] + new_date + text[match.end() :], True
    return text, False


def _is_standalone_date_label(text: str) -> bool:
    cleaned = text.strip().rstrip(":").strip()
    upper = cleaned.upper()
    if upper in ("DATE", "DATED", "INVOICE DATE", "NUMBER"):
        return True
    return bool(re.fullmatch(r"(?i)Invoice#\s*Date", cleaned))


def _paragraph_already_current(text: str, new_date: str) -> bool:
    stripped = text.strip()
    if not stripped:
        return True
    remainder = stripped.replace(new_date, "")
    return not DATE_VALUE.search(remainder)


def _update_inline_labeled_paragraph(paragraph, new_date: str) -> bool:
    raw = paragraph.text
    if _is_layout_only_text(raw):
        return False
    lead, text, trail = _split_outer_whitespace(raw)
    if not text:
        return False

    match = INLINE_DATE_LABELS.search(text)
    if match:
        if match.group(2).strip() == new_date:
            return False
        prefix = text[: match.start()]
        new_core = prefix + match.group(1) + new_date
        _replace_paragraph_text(paragraph, lead + new_core + trail)
        return True

    if _paragraph_already_current(text, new_date):
        return False

    if re.search(r"(?i)\bDate\s*:", text) and DATE_VALUE.search(text):
        new_core, changed = _replace_date_in_text(text, new_date)
        if changed:
            _replace_paragraph_text(paragraph, lead + new_core + trail)
            return True

    return False


def _update_paragraph_if_date_value(paragraph, new_date: str) -> bool:
    raw = paragraph.text
    if _is_layout_only_text(raw):
        return False
    lead, text, trail = _split_outer_whitespace(raw)
    if not text or _is_standalone_date_label(text):
        return False
    if _paragraph_already_current(text, new_date):
        return False
    if DATE_VALUE.search(text):
        new_core, changed = _replace_date_in_text(text, new_date)
        if changed:
            _replace_paragraph_text(paragraph, lead + new_core + trail)
            return True
    return False


def _iter_body_paragraphs(doc: Document):
    for paragraph in doc.paragraphs:
        yield paragraph


def _iter_cell_paragraphs(cell) -> list:
    """Paragraphs in a cell, including those inside nested tables."""
    paragraphs = list(cell.paragraphs)
    for nested in cell.tables:
        for row in nested.rows:
            for nested_cell in row.cells:
                paragraphs.extend(_iter_cell_paragraphs(nested_cell))
    return paragraphs


def _iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in _iter_cell_paragraphs(cell):
                yield paragraph


def _iter_paragraphs(doc: Document):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        yield from _iter_table_paragraphs(table)
    for section in doc.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            if header is not None:
                for paragraph in header.paragraphs:
                    yield paragraph
                for table in header.tables:
                    yield from _iter_table_paragraphs(table)
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            if footer is not None:
                for paragraph in footer.paragraphs:
                    yield paragraph
                for table in footer.tables:
                    yield from _iter_table_paragraphs(table)


def update_all_dates_in_document(
    doc: Document,
    invoice_date: date | None = None,
) -> int:
    """Update every Date / date / Dated field to the current date."""
    new_date = format_invoice_date(invoice_date)
    updated = 0
    seen: set[int] = set()

    def touch(paragraph) -> None:
        nonlocal updated
        pid = id(paragraph)
        if pid not in seen:
            seen.add(pid)
            updated += 1

    for paragraph in _iter_paragraphs(doc):
        if _is_layout_only_text(paragraph.text):
            continue
        if _update_inline_labeled_paragraph(paragraph, new_date):
            touch(paragraph)

    body_paragraphs = list(_iter_body_paragraphs(doc))
    for i, paragraph in enumerate(body_paragraphs):
        if not _is_standalone_date_label(paragraph.text):
            continue
        if i + 1 >= len(body_paragraphs):
            continue
        nxt = body_paragraphs[i + 1]
        if id(nxt) in seen:
            continue
        if _update_paragraph_if_date_value(nxt, new_date):
            touch(nxt)

    for paragraph in _iter_body_paragraphs(doc):
        if id(paragraph) in seen:
            continue
        if _is_layout_only_text(paragraph.text):
            continue
        if re.search(r"(?i)Invoice#\s*Date", paragraph.text):
            continue
        if _update_paragraph_if_date_value(paragraph, new_date):
            touch(paragraph)

    for table in doc.tables:
        for row in table.rows:
            row_has_date_header = any(
                p.text.strip().lower() == "date"
                for cell in row.cells
                for p in _iter_cell_paragraphs(cell)
            )
            if not row_has_date_header:
                continue
            for cell in row.cells:
                for paragraph in _iter_cell_paragraphs(cell):
                    if id(paragraph) in seen:
                        continue
                    if paragraph.text.strip().lower() == "date":
                        continue
                    if _update_paragraph_if_date_value(paragraph, new_date):
                        touch(paragraph)

    return updated


def update_invoice3_dates(
    doc: Document,
    invoice_date: date | None = None,
) -> None:
    """Invoice 3: update only the date value + Pay by; preserve header label layout."""
    new_date = format_invoice_date(invoice_date)
    updated = 0

    if len(doc.paragraphs) > 12:
        paragraph = doc.paragraphs[12]
        if not _is_layout_only_text(paragraph.text):
            lead, core, trail = _split_outer_whitespace(paragraph.text)
            if DATE_VALUE.search(core):
                new_core, changed = _replace_date_in_text(core, new_date)
                if changed:
                    _replace_paragraph_text(paragraph, lead + new_core + trail)
                    updated += 1

    if len(doc.tables) > 3:
        cell = doc.tables[3].rows[0].cells[0]
        for paragraph in cell.paragraphs:
            if _is_layout_only_text(paragraph.text):
                continue
            if _update_inline_labeled_paragraph(paragraph, new_date):
                updated += 1

    if updated == 0:
        raise ValidationError(
            "Could not find a Date or Pay by field to update in invoice template 3."
        )


def update_dated_in_document(
    doc: Document,
    invoice_date: date | None = None,
    invoice_number: int | None = None,
) -> None:
    if invoice_number == 3:
        update_invoice3_dates(doc, invoice_date)
        return
    count = update_all_dates_in_document(doc, invoice_date)
    if count == 0:
        # New company templates (6+) often lack a standard Date label; skip rather than fail.
        if invoice_number is not None and invoice_number >= 6:
            return
        label = f"invoice template {invoice_number}" if invoice_number else "document"
        raise ValidationError(
            f'Could not find a Date, Invoice Date, or Pay by field to update in {label}.'
        )


def copy_and_update_date(
    source_path: Path,
    output_path: Path,
    invoice_date: date | None = None,
    invoice_number: int | None = None,
) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source_path, output_path)
    doc = Document(str(output_path))
    update_dated_in_document(doc, invoice_date, invoice_number)
    doc.save(str(output_path))
    return output_path

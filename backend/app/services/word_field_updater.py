import re

from docx import Document

from backend.app.models import InvoiceTask, WordInvoiceData
from backend.app.services.amount_distribution import distribute_amounts
from backend.app.services.word_date_updater import update_dated_in_document
from backend.app.services.word_template_config import LineItemConfig, line_item_config

NAME_LINE_PATTERN = re.compile(r"(?i)^Name:\s*(.*)$")
ACCOUNT_NAME_PATTERN = re.compile(
    r"(?i)Account Name:\s*(.+?)(?=\s+Account No\.?:|\s+Account\s+No\.?:|$)"
)
SEND_PAYMENT_TO_PATTERN = re.compile(r"(?i)^(Send Payment To\s*:\s*)(.+)$")
BANK_LINE_PATTERN = re.compile(r"(?i)^(Bank Name\s*:|Bank\s*:)(.*)$")
BANK_NAME_MULTILINE_PATTERN = re.compile(
    r"(?i)^Bank\s*:(.*?)\n\s*Name\s*:(.*)$",
    re.DOTALL,
)
IBAN_LINE_PATTERN = re.compile(r"(?i)^(IBAN NO\.?\s*:|IBAN\s*:)(.*)$")
NAME_IBAN_BLOCK_PATTERN = re.compile(
    r"(?i)^(Name\s*:)([^\n]*)(\n\s*IBAN(?:\s*NO\.?)?\s*:)([^\n]*)$",
    re.DOTALL,
)
BRANCH_LINE_PATTERN = re.compile(r"(?i)^(Branch Code\s*:)(.*)$")
BRANCH_IN_TEXT_PATTERN = re.compile(r"(?i)(Branch Code\s*:)\s*[^\n\r]*")
IBAN_BRANCH_BLOCK_PATTERN = re.compile(
    r"(?i)^(IBAN(?:\s*NO\.?)?\s*:)([^\n]*)(\n\s*Branch Code\s*:)([^\n]*)"
)
ACCOUNT_NAME_LINE_PATTERN = re.compile(r"(?i)^(Account Name\s*:)(.*)$")
PAYMENT_NAME_LINE_PATTERN = re.compile(r"(?i)^(Name\s*:)(.*)$")
INVOICE_NO_IN_TEXT = re.compile(r"(?i)(Invoice No\.?\s*:)\s*[^\n\r]*")
TOTAL_AMOUNT_PATTERN = re.compile(r"(\$[\d,]+(?:\.\d{2})?)")


def format_amount(
    amount: int,
    *,
    cents: bool = False,
    compact: bool | None = None,
) -> str:
    """
    Format money for invoice templates.

    Amounts >= 1000 omit thousands separators by default so narrow TOTAL /
    amount cells do not wrap or push a second PDF page
    (e.g. $1500 / $1500.00 instead of $1,500 / $1,500.00).
    Pass compact=False to force comma grouping.
    """
    value = int(amount)
    if compact is None:
        compact = abs(value) >= 1_000
    if cents:
        # Very large + cents overflows narrow TOTAL cells ($250000.00 → wrap).
        if compact and abs(value) >= 10_000:
            return f"${value}"
        return f"${value:.2f}" if compact else f"${value:,.2f}"
    return f"${value}" if compact else f"${value:,}"


def _rebuild_bank_account_block(text: str, data: WordInvoiceData) -> str | None:
    """
    Rebuild payment paragraphs that use Bank Name / Account Name / Account No / Branch Code.
    Keeps any prefix before the first Bank Name label (e.g. company + PAY TO:).
    """
    match = re.search(r"(?i)bank\s*name", text)
    if not match:
        return None
    prefix = text[: match.start()]
    bank = data.bank or ""
    name = data.person_name or ""
    iban = data.iban or ""
    branch = data.branch_code or ""
    # Prefer newline style when original had newlines.
    if "\n" in text[match.start() :]:
        body = (
            f"Bank Name: {bank}\n"
            f"Account Name: {name}\n"
            f"Account No.: {iban}\n"
            f"Branch Code: {branch}"
        )
    else:
        body = (
            f"Bank Name: {bank} Account Name: {name} "
            f"Account No.: {iban} Branch Code: {branch}"
        )
    return prefix + body


def _is_layout_only_text(text: str) -> bool:
    return not text.strip() or bool(re.fullmatch(r"[\s\t]+", text))


def _split_outer_whitespace(text: str) -> tuple[str, str, str]:
    core = text.strip()
    if not core:
        return "", text, ""
    start = text.index(core)
    end = start + len(core)
    return text[:start], core, text[end:]


def _replace_paragraph_text(paragraph, new_text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)


def _iter_cell_paragraphs(cell):
    paragraphs = list(cell.paragraphs)
    for nested in cell.tables:
        for row in nested.rows:
            for nested_cell in row.cells:
                paragraphs.extend(_iter_cell_paragraphs(nested_cell))
    return paragraphs


def _iter_all_paragraphs(doc: Document):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in _iter_cell_paragraphs(cell):
                    yield paragraph


def _update_branch_code_in_paragraph(paragraph, branch: str) -> bool:
    if not branch or _is_layout_only_text(paragraph.text):
        return False
    text = paragraph.text
    if not re.search(r"(?i)Branch\s*Code\s*:", text):
        return False
    new_text = BRANCH_IN_TEXT_PATTERN.sub(
        lambda match: f"{match.group(1)} {branch}", text, count=1
    )
    if new_text == text:
        return False
    _replace_paragraph_text(paragraph, new_text)
    return True


def _update_payment_line(paragraph, pattern: re.Pattern, value: str) -> bool:
    if not value or _is_layout_only_text(paragraph.text):
        return False
    lead, core, trail = _split_outer_whitespace(paragraph.text)
    match = pattern.match(core)
    if not match:
        return False
    prefix = match.group(1)
    sep = "" if prefix.endswith((" ", ":")) else " "
    if prefix.rstrip().endswith(":") and not prefix.endswith(": "):
        sep = " "
    new_core = f"{prefix}{sep}{value}".replace(":  ", ": ")
    _replace_paragraph_text(paragraph, f"{lead}{new_core}{trail}")
    return True


def _replace_labeled_line_value(paragraph, label: str, new_value: str) -> bool:
    """Replace only the value on a single labeled line; keep surrounding whitespace."""
    if _is_layout_only_text(paragraph.text):
        return False
    lead, core, trail = _split_outer_whitespace(paragraph.text)
    pattern = re.compile(rf"(?i)^({re.escape(label)}\s*:)( *)(.*)$")
    match = pattern.match(core)
    if not match:
        return False
    new_core = f"{match.group(1)}{match.group(2)}{new_value}"
    _replace_paragraph_text(paragraph, f"{lead}{new_core}{trail}")
    return True


def _update_invoice_no_fields(doc: Document, invoice_no: str) -> None:
    """Fill Invoice No. in the document with the incremented value from the sheet."""
    if not invoice_no:
        return
    for paragraph in _iter_all_paragraphs(doc):
        if _is_layout_only_text(paragraph.text):
            continue
        text = paragraph.text
        if _replace_labeled_line_value(paragraph, "Invoice No.", invoice_no):
            continue
        if _replace_labeled_line_value(paragraph, "Invoice Number", invoice_no):
            continue
        if _replace_labeled_line_value(paragraph, "INVOICE NO", invoice_no):
            continue
        # "INVOICE NO. 000001" / "NO. 000001" style without a blank value slot
        if re.search(r"(?i)invoice\s*no\.?\s*\d", text) or re.match(
            r"(?i)^NO\.\s*\d", text.strip()
        ):
            new_text = re.sub(
                r"(?i)(invoice\s*no\.?\s*|NO\.\s*)(\d[\w\-]*)",
                lambda m: f"{m.group(1)}{invoice_no}",
                text,
                count=1,
            )
            if new_text != text:
                _replace_paragraph_text(paragraph, new_text)
                continue
        if not re.search(r"(?i)Invoice No|Invoice Number|Invoice#", text):
            continue
        new_text = INVOICE_NO_IN_TEXT.sub(
            lambda match: f"{match.group(1)} {invoice_no}", paragraph.text, count=1
        )
        if new_text != paragraph.text:
            _replace_paragraph_text(paragraph, new_text)

    # Table cells: Invoice# value in adjacent cell (Alpha Digital etc.)
    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            for ci, cell in enumerate(cells):
                label = cell.text.strip().lower().replace(" ", "")
                if label in ("invoice#", "invoicenr:", "invoicenr") and ci + 1 < len(
                    cells
                ):
                    _set_cell_text(cells[ci + 1], invoice_no)


def _update_payment_fields(doc: Document, data: WordInvoiceData) -> None:
    """Fill Bank, Name, IBAN, and Branch Code from payment sheet data."""
    inv = data.invoice_number
    name = data.person_name
    bank = data.bank
    iban = data.iban
    branch = data.branch_code or ""

    for paragraph in _iter_all_paragraphs(doc):
        if _is_layout_only_text(paragraph.text):
            continue
        text = paragraph.text.strip()
        if not text:
            continue

        if branch:
            iban_branch = IBAN_BRANCH_BLOCK_PATTERN.match(text)
            if iban_branch:
                _replace_paragraph_text(
                    paragraph,
                    f"{iban_branch.group(1)} {iban}"
                    f"{iban_branch.group(3)} {branch}",
                )
                continue

        # Invoice 2 (and similar): Name and IBAN share one paragraph with a newline.
        name_iban = NAME_IBAN_BLOCK_PATTERN.match(text)
        if name_iban:
            name_label = name_iban.group(1)
            iban_label = name_iban.group(3)
            _replace_paragraph_text(
                paragraph,
                f"{name_label} {name}{iban_label} {iban}".replace(":  ", ": "),
            )
            continue

        if re.search(r"(?i)Bank\s*:", text) and re.search(
            r"(?i)Account Name:", text
        ):
            _replace_paragraph_text(
                paragraph, f"Bank: {bank} Account Name: {name}"
            )
            continue

        if BANK_NAME_MULTILINE_PATTERN.match(text.strip()):
            _replace_paragraph_text(paragraph, f"Bank: {bank}\nName: {name}")
            continue

        if _update_payment_line(paragraph, BANK_LINE_PATTERN, bank):
            continue
        if _update_payment_line(paragraph, IBAN_LINE_PATTERN, iban):
            continue
        if re.match(r"(?i)^Account\s*No\.?\s*:", text):
            _replace_labeled_line_value(paragraph, "Account No.", iban)
            continue
        if re.match(r"(?i)^Account\s*NO\s*:", text):
            _replace_labeled_line_value(paragraph, "Account NO", iban)
            continue
        if branch and _update_payment_line(paragraph, BRANCH_LINE_PATTERN, branch):
            continue

        if _update_payment_line(paragraph, ACCOUNT_NAME_LINE_PATTERN, f" {name}".strip()):
            continue

        # Multiline Bank Name / Account Name / Account No / Branch Code blocks
        if re.search(r"(?i)bank\s*name", text) and re.search(
            r"(?i)account\s*name", text
        ):
            rebuilt = _rebuild_bank_account_block(text, data)
            if rebuilt and rebuilt != paragraph.text:
                _replace_paragraph_text(paragraph, rebuilt)
                continue

        if PAYMENT_NAME_LINE_PATTERN.match(text) and not re.search(
            r"(?i)account\s+name|bank\s+name", text
        ):
            _replace_labeled_line_value(paragraph, "Name", name)
            continue

        match = SEND_PAYMENT_TO_PATTERN.match(text)
        if match:
            _replace_paragraph_text(paragraph, match.group(1) + name)
            if inv == 1:
                _update_signature_name_before_dated(doc, name)
            continue

        match = ACCOUNT_NAME_PATTERN.search(paragraph.text)
        if match and bank.lower() in paragraph.text.lower():
            prefix = paragraph.text[: match.start()]
            suffix = paragraph.text[match.end() :]
            bank_prefix = ""
            bank_match = re.search(r"(?i)Bank\s*:\s*[^A]+", paragraph.text)
            if bank_match:
                bank_prefix = f"Bank: {bank} "
            _replace_paragraph_text(
                paragraph,
                f"{bank_prefix}Account Name: {name}{suffix}",
            )
            continue

        if re.match(r"(?i)^Bank\s*:", text) and "Account Name" not in text:
            _replace_labeled_line_value(paragraph, "Bank", bank)
            continue
        if re.match(r"(?i)^IBAN", text) and "branch code" not in text.lower():
            label = "IBAN No." if "NO." in text.upper() or inv in (4, 5) else "IBAN"
            _replace_labeled_line_value(paragraph, label, iban)

    for paragraph in _iter_all_paragraphs(doc):
        _update_branch_code_in_paragraph(paragraph, branch)


def _update_name_field(doc: Document, person_name: str) -> None:
    """Update Name / Send Payment To / signature name fields."""
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if NAME_LINE_PATTERN.match(text):
            _replace_paragraph_text(paragraph, f"Name: {person_name}")
            return
        match = SEND_PAYMENT_TO_PATTERN.match(text)
        if match:
            _replace_paragraph_text(paragraph, match.group(1) + person_name)
            _update_signature_name_before_dated(doc, person_name)
            return
        match = ACCOUNT_NAME_PATTERN.search(paragraph.text)
        if match:
            prefix = paragraph.text[: match.start()]
            suffix = paragraph.text[match.end() :]
            _replace_paragraph_text(
                paragraph, f"{prefix}Account Name: {person_name}{suffix}"
            )
            return

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if NAME_LINE_PATTERN.match(paragraph.text.strip()):
                        _replace_paragraph_text(paragraph, f"Name: {person_name}")
                        return


def _update_signature_name_before_dated(doc: Document, person_name: str) -> None:
    """Invoice 1: name line above bottom DATED: (all caps, not top Dated:)."""
    for i, paragraph in enumerate(doc.paragraphs):
        if not paragraph.text.strip().startswith("DATED:"):
            continue
        for j in range(i - 1, -1, -1):
            prev = doc.paragraphs[j]
            prev_text = prev.text.strip()
            if not prev_text:
                continue
            if ":" not in prev_text and len(prev_text) < 80:
                _replace_paragraph_text(prev, person_name)
            break
        break


def _compact_label(text: str) -> str:
    return re.sub(r"[\s:]", "", text.strip().lower())


def _is_sub_total_label(text: str) -> bool:
    return "subtotal" in _compact_label(text)


def _is_total_only_label(text: str) -> bool:
    if not text or not text.strip():
        return False
    if _is_sub_total_label(text):
        return False
    compact = _compact_label(text)
    return compact in ("total", "tota")


def _replace_amount_in_paragraph(paragraph, amount_str: str) -> bool:
    """Update only $ amount run text so labels/spacing keep their original font."""
    updated = False
    for run in paragraph.runs:
        if TOTAL_AMOUNT_PATTERN.search(run.text):
            run.text = TOTAL_AMOUNT_PATTERN.sub(amount_str, run.text, count=1)
            updated = True
    return updated


def _set_cell_text(cell, text: str) -> None:
    if not cell.paragraphs:
        cell.add_paragraph(text)
        return
    non_empty = [p for p in cell.paragraphs if p.text.strip()]
    targets = non_empty if non_empty else [cell.paragraphs[-1]]
    for para in targets:
        _replace_paragraph_text(para, text)


def _set_cell_all_amounts(cell, total_amount: int) -> None:
    amount_str = format_amount(total_amount)
    if not cell.paragraphs:
        cell.add_paragraph(amount_str)
        return
    non_empty = [p for p in cell.paragraphs if p.text.strip()]
    targets = non_empty if non_empty else [cell.paragraphs[-1]]
    for para in targets:
        _replace_paragraph_text(para, amount_str)


def _amount_cell(row, amount_col: int):
    if len(row.cells) > amount_col:
        return row.cells[amount_col]
    return row.cells[-1]


def _row_has_sub_total(row, amount_col: int) -> bool:
    for ci, cell in enumerate(row.cells):
        if ci >= amount_col:
            break
        if _is_sub_total_label(cell.text):
            return True
        for para in cell.paragraphs:
            if _is_sub_total_label(para.text):
                return True
    return False


def _row_has_total_only(row, amount_col: int) -> bool:
    for ci, cell in enumerate(row.cells):
        if ci >= amount_col:
            break
        if _is_total_only_label(cell.text):
            return True
        for para in cell.paragraphs:
            if _is_total_only_label(para.text):
                return True
    return False


def _is_tax_label(text: str) -> bool:
    return "tax" in _compact_label(text)


def _row_has_tax(row, amount_col: int) -> bool:
    for ci, cell in enumerate(row.cells):
        if ci >= amount_col:
            break
        if _is_tax_label(cell.text):
            return True
        for para in cell.paragraphs:
            if _is_tax_label(para.text):
                return True
    return False


def _ensure_tax_zero(table, amount_col: int) -> None:
    for row in table.rows:
        if _row_has_tax(row, amount_col):
            _set_cell_all_amounts(_amount_cell(row, amount_col), 0)


def _update_subtotal_and_total(table, total_amount: int, amount_col: int) -> None:
    """Sub Total and T O T A L / Total rows in table get the same frontend total."""
    sub_total_rows: list[int] = []

    for ri, row in enumerate(table.rows):
        if len(row.cells) <= amount_col:
            continue

        amount_cell = _amount_cell(row, amount_col)

        for ci in range(min(amount_col, len(row.cells))):
            cell = row.cells[ci]
            label_paras = [p for p in cell.paragraphs if p.text.strip()]
            if len(label_paras) < 2:
                continue
            has_sub = any(_is_sub_total_label(p.text) for p in label_paras)
            has_total = any(_is_total_only_label(p.text) for p in label_paras)
            if has_sub:
                sub_total_rows.append(ri)
            if has_sub or has_total:
                _set_cell_all_amounts(amount_cell, total_amount)
                while len(amount_cell.paragraphs) < len(label_paras):
                    amount_cell.add_paragraph(format_amount(total_amount))
                for para in amount_cell.paragraphs[: len(label_paras)]:
                    _replace_paragraph_text(para, format_amount(total_amount))

        if _row_has_sub_total(row, amount_col):
            sub_total_rows.append(ri)
            _set_cell_all_amounts(amount_cell, total_amount)
        elif _row_has_total_only(row, amount_col):
            _set_cell_all_amounts(amount_cell, total_amount)

    for ri in set(sub_total_rows):
        next_ri = ri + 1
        if next_ri < len(table.rows):
            next_row = table.rows[next_ri]
            if _row_has_tax(next_row, amount_col):
                continue
            if _row_has_total_only(next_row, amount_col) and len(next_row.cells) > amount_col:
                _set_cell_all_amounts(_amount_cell(next_row, amount_col), total_amount)

    for row in table.rows:
        if _row_has_total_only(row, amount_col) and len(row.cells) > amount_col:
            _set_cell_all_amounts(_amount_cell(row, amount_col), total_amount)


def _update_paragraph_totals(doc: Document, total_amount: int) -> None:
    """Update Total lines outside the table (e.g. Invoice 5 paragraph Total)."""
    amount_str = format_amount(total_amount)
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text.lower().startswith("total"):
            continue
        if _is_sub_total_label(text):
            continue
        if "\t" in text:
            _replace_paragraph_text(paragraph, f"Total\t{amount_str}")
        elif TOTAL_AMOUNT_PATTERN.search(text):
            new_text = TOTAL_AMOUNT_PATTERN.sub(amount_str, text, count=1)
            _replace_paragraph_text(paragraph, new_text)
        elif text.lower() in ("total", "total:"):
            _replace_paragraph_text(paragraph, f"Total\t{amount_str}")


def _update_line_items(
    table,
    tasks: list[InvoiceTask],
    config: LineItemConfig,
) -> None:
    for i, row_idx in enumerate(config.line_item_rows):
        if i >= len(tasks) or row_idx >= len(table.rows):
            continue
        row = table.rows[row_idx]
        task = tasks[i]
        if len(row.cells) > config.desc_col:
            _set_cell_text(row.cells[config.desc_col], task.sub_task)
        if len(row.cells) > config.amount_col:
            _set_cell_all_amounts(row.cells[config.amount_col], task.amount)
    for row_idx in config.clear_rows:
        if row_idx >= len(table.rows):
            continue
        row = table.rows[row_idx]
        if len(row.cells) > config.desc_col:
            _set_cell_text(row.cells[config.desc_col], "")
        if len(row.cells) > config.amount_col:
            _set_cell_text(row.cells[config.amount_col], "")


def _all_document_tables(doc: Document) -> list:
    """All tables in the body, including those nested in text boxes."""
    from docx.oxml.ns import qn
    from docx.table import Table

    tables: list = []
    seen: set[int] = set()
    for tbl in doc.element.body.findall(".//" + qn("w:tbl")):
        key = id(tbl)
        if key in seen:
            continue
        seen.add(key)
        tables.append(Table(tbl, doc))
    return tables


def _resolve_table(doc: Document, config: LineItemConfig):
    tables = _all_document_tables(doc) if config.use_all_tables else list(doc.tables)
    if config.table_index >= len(tables):
        return None
    return tables[config.table_index]


def _set_paragraph_amount_by_label(doc: Document, labels: tuple[str, ...], amount: int) -> None:
    amount_str = format_amount(amount)
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        compact = _compact_label(text)
        if any(_compact_label(label) == compact or compact.startswith(_compact_label(label)) for label in labels):
            if TOTAL_AMOUNT_PATTERN.search(paragraph.text):
                new_text = TOTAL_AMOUNT_PATTERN.sub(amount_str, paragraph.text, count=1)
                _replace_paragraph_text(paragraph, new_text)
            elif text.upper().startswith("TOTAL") and "\t" in paragraph.text:
                _update_labeled_tab_amount(paragraph, "Total", amount)


def _fill_standard_club_template(doc: Document, data: WordInvoiceData) -> None:
    """Shared fill for new templates 6–13: payment + 3 club lines + totals."""
    config = line_item_config(data.invoice_number)
    if data.bank or data.iban or data.branch_code or data.person_name:
        _update_payment_fields(doc, data)
    table = _resolve_table(doc, config)
    if table is not None and data.tasks and len(data.tasks) >= 3:
        _update_line_items(table, data.tasks[:3], config)
        _update_subtotal_and_total(table, data.total_amount, config.amount_col)
    _update_paragraph_totals(doc, data.total_amount)


def _apply_template_invoice_logic(
    doc: Document,
    data: WordInvoiceData,
    config: LineItemConfig,
) -> None:
    _update_name_field(doc, data.person_name)
    if data.bank or data.iban or data.branch_code:
        _update_payment_fields(doc, data)

    if len(doc.tables) > config.table_index:
        table = doc.tables[config.table_index]
        if data.tasks and len(data.tasks) >= 3:
            _update_line_items(table, data.tasks[:3], config)
        else:
            amounts = distribute_amounts(data.total_amount)
            for i, row_idx in enumerate(config.line_item_rows):
                if row_idx >= len(table.rows):
                    continue
                row = table.rows[row_idx]
                if len(row.cells) > config.amount_col:
                    _set_cell_all_amounts(row.cells[config.amount_col], amounts[i])
        _update_subtotal_and_total(table, data.total_amount, config.amount_col)

    _update_paragraph_totals(doc, data.total_amount)


def _apply_template_amount_logic(
    doc: Document,
    data: WordInvoiceData,
    amount_col: int,
    table_index: int = 0,
    line_item_rows: list[int] | None = None,
    desc_col: int = 0,
) -> None:
    config = LineItemConfig(
        table_index=table_index,
        line_item_rows=line_item_rows or [1, 2, 3],
        desc_col=desc_col,
        amount_col=amount_col,
    )
    _apply_template_invoice_logic(doc, data, config)


def _update_labeled_tab_amount(paragraph, label: str, amount: int) -> None:
    if _is_layout_only_text(paragraph.text):
        return
    lead, core, trail = _split_outer_whitespace(paragraph.text)
    amount_str = format_amount(amount)
    if "\t" in core:
        new_core = re.sub(
            r"(\$[\d,]+(?:\.\d{2})?)",
            amount_str,
            core,
            count=1,
        )
        if new_core == core:
            new_core = f"{label}\t{amount_str}"
    else:
        new_core = f"{label}\t{amount_str}"
    _replace_paragraph_text(paragraph, f"{lead}{new_core}{trail}")


def _update_invoice3_send_payment_to(doc: Document, person_name: str) -> None:
    for paragraph in doc.paragraphs:
        if _is_layout_only_text(paragraph.text):
            continue
        lead, core, trail = _split_outer_whitespace(paragraph.text)
        match = SEND_PAYMENT_TO_PATTERN.match(core)
        if not match:
            continue
        _replace_paragraph_text(
            paragraph, f"{lead}{match.group(1)}{person_name}{trail}"
        )
        return


def _find_invoice3_summary_table(doc: Document):
    """Locate the SUBTOTAL/TAX/TOTAL summary table (index differs for HTML vs Word)."""
    for table in doc.tables:
        if not table.rows or not table.rows[0].cells:
            continue
        blob = table.rows[0].cells[0].text.upper()
        if "SUBTOTAL" in blob and "TOTAL" in blob:
            return table
    return None


def _update_invoice3_summary_table(doc: Document, total_amount: int) -> None:
    """Invoice 3: SUBTOTAL and TOTAL = total, TAX always $0."""
    table = _find_invoice3_summary_table(doc)
    if table is None:
        return
    cell = table.rows[0].cells[0]
    for para in cell.paragraphs:
        if _is_layout_only_text(para.text):
            continue
        text = para.text.strip().upper()
        if text.startswith("SUBTOTAL"):
            _update_labeled_tab_amount(para, "SUBTOTAL:", total_amount)
        elif text.startswith("TAX"):
            _update_labeled_tab_amount(para, "TAX:", 0)
        elif text.startswith("TOTAL"):
            _update_labeled_tab_amount(para, "TOTAL:", total_amount)


def _update_invoice4_grand_total(doc: Document, total_amount: int) -> None:
    """Invoice 4: grand Total lives in table 3, separate from the line-item table."""
    if len(doc.tables) < 4:
        return
    amount_str = format_amount(total_amount)
    cell = doc.tables[3].rows[0].cells[0]
    for para in cell.paragraphs:
        if _replace_amount_in_paragraph(para, amount_str):
            continue
        if re.search(r"(?i)total", para.text) and TOTAL_AMOUNT_PATTERN.search(para.text):
            new_text = TOTAL_AMOUNT_PATTERN.sub(amount_str, para.text, count=1)
            if new_text != para.text:
                _replace_paragraph_text(para, new_text)


def _update_template_1(doc: Document, data: WordInvoiceData) -> None:
    _apply_template_invoice_logic(doc, data, line_item_config(1))


def _update_template_2(doc: Document, data: WordInvoiceData) -> None:
    config = line_item_config(2)
    _apply_template_invoice_logic(doc, data, config)
    if len(doc.tables) > config.table_index:
        _ensure_tax_zero(doc.tables[config.table_index], amount_col=config.amount_col)


def _update_template_3(doc: Document, data: WordInvoiceData) -> None:
    """Invoice 3: update payment, items, summary, and name only; preserve header layout."""
    config = line_item_config(3)
    _update_invoice3_send_payment_to(doc, data.person_name)
    if data.bank or data.iban or data.branch_code:
        _update_payment_fields(doc, data)
    if len(doc.tables) > config.table_index:
        table = doc.tables[config.table_index]
        if data.tasks and len(data.tasks) >= 3:
            _update_line_items(table, data.tasks[:3], config)
        else:
            amounts = distribute_amounts(data.total_amount)
            for i, row_idx in enumerate(config.line_item_rows):
                if row_idx >= len(table.rows):
                    continue
                row = table.rows[row_idx]
                if len(row.cells) > config.amount_col:
                    _set_cell_all_amounts(row.cells[config.amount_col], amounts[i])
    _update_invoice3_summary_table(doc, data.total_amount)


def _update_template_4(doc: Document, data: WordInvoiceData) -> None:
    """Invoice 4: update payment table + line-item table only; preserve header layout."""
    config = line_item_config(4)
    if data.bank or data.iban or data.branch_code:
        _update_payment_fields(doc, data)
    if len(doc.tables) > config.table_index:
        table = doc.tables[config.table_index]
        if data.tasks and len(data.tasks) >= 3:
            _update_line_items(table, data.tasks[:3], config)
        else:
            amounts = distribute_amounts(data.total_amount)
            for i, row_idx in enumerate(config.line_item_rows):
                if row_idx >= len(table.rows):
                    continue
                row = table.rows[row_idx]
                if len(row.cells) > config.amount_col:
                    _set_cell_all_amounts(row.cells[config.amount_col], amounts[i])
        _update_subtotal_and_total(table, data.total_amount, config.amount_col)
    _update_invoice4_grand_total(doc, data.total_amount)


def _update_template_5(doc: Document, data: WordInvoiceData) -> None:
    _apply_template_invoice_logic(doc, data, line_item_config(5))


def _update_template_6_ravotek(doc: Document, data: WordInvoiceData) -> None:
    _fill_standard_club_template(doc, data)
    if len(doc.tables) > 1:
        cell = doc.tables[1].rows[0].cells[-1]
        _set_cell_all_amounts(cell, data.total_amount)


def _update_template_7_ignitai(doc: Document, data: WordInvoiceData) -> None:
    _fill_standard_club_template(doc, data)
    tables = _all_document_tables(doc)
    if tables:
        table = tables[0]
        # Total Amount Due row
        if len(table.rows) > 8 and len(table.rows[8].cells) > 3:
            _set_cell_all_amounts(table.rows[8].cells[3], data.total_amount)
    # Invoice number sits on the paragraph after "Invoice Number:"
    if data.document_invoice_no:
        for i, paragraph in enumerate(doc.paragraphs):
            if paragraph.text.strip().lower().startswith("invoice number"):
                if i + 1 < len(doc.paragraphs):
                    _replace_paragraph_text(
                        doc.paragraphs[i + 1], data.document_invoice_no
                    )
                break


def _update_template_8_coretechify(doc: Document, data: WordInvoiceData) -> None:
    _fill_standard_club_template(doc, data)
    table = doc.tables[0] if doc.tables else None
    if table and len(table.rows) > 6:
        _set_cell_all_amounts(table.rows[6].cells[-1], data.total_amount)


def _update_template_9_ecomify(doc: Document, data: WordInvoiceData) -> None:
    _fill_standard_club_template(doc, data)
    if len(doc.tables) > 2:
        # subtotal / total amount rows
        summary = doc.tables[2]
        if len(summary.rows) > 0:
            _set_cell_all_amounts(summary.rows[0].cells[-1], data.total_amount)
        if len(summary.rows) > 3:
            _set_cell_all_amounts(summary.rows[3].cells[-1], data.total_amount)
        # zero discount/tax for clean totals
        if len(summary.rows) > 1:
            _set_cell_text(summary.rows[1].cells[-1], "$0")
        if len(summary.rows) > 2:
            _set_cell_text(summary.rows[2].cells[-1], "$0")
    if data.document_invoice_no and doc.tables:
        inv_table = doc.tables[0]
        if inv_table.rows and len(inv_table.rows[0].cells) > 1:
            _set_cell_text(inv_table.rows[0].cells[1], data.document_invoice_no)


def _update_template_10_cozy(doc: Document, data: WordInvoiceData) -> None:
    _fill_standard_club_template(doc, data)
    amount_str = format_amount(data.total_amount)
    for i, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip().upper() == "TOTAL" and i + 1 < len(doc.paragraphs):
            _replace_paragraph_text(doc.paragraphs[i + 1], amount_str)
            break


def _update_template_11_beecodify(doc: Document, data: WordInvoiceData) -> None:
    _fill_standard_club_template(doc, data)
    table = doc.tables[0] if doc.tables else None
    if table and len(table.rows) > 6:
        _set_cell_all_amounts(table.rows[6].cells[-1], data.total_amount)
    if data.document_invoice_no:
        for paragraph in doc.paragraphs:
            if re.match(r"(?i)^NO\.\s*", paragraph.text.strip()):
                _replace_paragraph_text(paragraph, f"NO. {data.document_invoice_no}")
                break


def _update_template_12_bravix(doc: Document, data: WordInvoiceData) -> None:
    _fill_standard_club_template(doc, data)
    table = doc.tables[0] if doc.tables else None
    if table and len(table.rows) > 6:
        row = table.rows[6]
        if row.cells:
            _set_cell_text(
                row.cells[0],
                (
                    f"Account Name: {data.person_name}\n"
                    f"Bank Name: {data.bank}\n"
                    f"Account No: {data.iban}\n"
                    f"Branch Code: {data.branch_code or ''}"
                ),
            )
        _set_cell_all_amounts(row.cells[-1], data.total_amount)
        if len(row.cells) > 4:
            _set_cell_all_amounts(row.cells[4], data.total_amount)


def _update_template_13_alpha(doc: Document, data: WordInvoiceData) -> None:
    _fill_standard_club_template(doc, data)
    amount_str = format_amount(data.total_amount)
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().lower().startswith("total"):
            if "\t" in paragraph.text or TOTAL_AMOUNT_PATTERN.search(paragraph.text):
                _update_labeled_tab_amount(paragraph, "Total", data.total_amount)
            else:
                _replace_paragraph_text(paragraph, f"Total\t{amount_str}")
    # Payment lines are split across paragraphs
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if re.match(r"(?i)^bank\s*name\s*:", text):
            _replace_paragraph_text(paragraph, f"Bank Name: {data.bank}")
        elif re.match(r"(?i)^account\s*name\s*:", text):
            _replace_paragraph_text(
                paragraph,
                (
                    f"Account Name: {data.person_name} "
                    f"Account No.: {data.iban} "
                    f"Branch Code: {data.branch_code or ''}"
                ),
            )


def _update_template_14_synergo(doc: Document, data: WordInvoiceData) -> None:
    if data.bank or data.iban or data.branch_code or data.person_name:
        _update_payment_fields(doc, data)
    if not doc.tables or not data.tasks or len(data.tasks) < 3:
        return
    table = doc.tables[0]
    # Title rows 1/4/6 hold the three club items; detail rows get amounts.
    title_rows = [1, 4, 6]
    detail_rows = [2, 5, 7]
    for i, (title_i, detail_i) in enumerate(zip(title_rows, detail_rows)):
        task = data.tasks[i]
        if title_i < len(table.rows) and table.rows[title_i].cells:
            _set_cell_text(table.rows[title_i].cells[0], task.sub_task)
        if detail_i < len(table.rows):
            row = table.rows[detail_i]
            if row.cells:
                _set_cell_text(row.cells[0], "")
            if len(row.cells) > 3:
                _set_cell_all_amounts(row.cells[3], task.amount)
    # Clear decorative middle row
    if len(table.rows) > 3:
        for cell in table.rows[3].cells:
            _set_cell_text(cell, "")
    if len(doc.tables) > 1 and doc.tables[1].rows:
        _set_cell_all_amounts(doc.tables[1].rows[0].cells[-1], data.total_amount)
    if data.document_invoice_no:
        for i, paragraph in enumerate(doc.paragraphs):
            if "invoice number" in paragraph.text.lower():
                # value often on next paragraph
                if i + 2 < len(doc.paragraphs) and doc.paragraphs[i + 2].text.strip().isdigit():
                    _replace_paragraph_text(
                        doc.paragraphs[i + 2], data.document_invoice_no
                    )
                elif not paragraph.text.strip().endswith((":", ": ")):
                    _replace_labeled_line_value(
                        paragraph, "Invoice Number", data.document_invoice_no
                    )
                else:
                    # keep label; put value on following numeric para if present
                    for j in range(i + 1, min(i + 4, len(doc.paragraphs))):
                        if doc.paragraphs[j].text.strip():
                            _replace_paragraph_text(
                                doc.paragraphs[j], data.document_invoice_no
                            )
                            break
                break


def update_invoice_fields(doc: Document, data: WordInvoiceData) -> None:
    handlers = {
        1: _update_template_1,
        2: _update_template_2,
        3: _update_template_3,
        4: _update_template_4,
        5: _update_template_5,
        6: _update_template_6_ravotek,
        7: _update_template_7_ignitai,
        8: _update_template_8_coretechify,
        9: _update_template_9_ecomify,
        10: _update_template_10_cozy,
        11: _update_template_11_beecodify,
        12: _update_template_12_bravix,
        13: _update_template_13_alpha,
        14: _update_template_14_synergo,
    }
    handler = handlers.get(data.invoice_number)
    if handler:
        handler(doc, data)
    # Some templates keep the invoice value on the next paragraph / own cell;
    # generic label rewrite can corrupt textbox layouts (e.g. Ignitai).
    skip_generic_invoice_no = {7, 9, 11, 14}
    if data.document_invoice_no and data.invoice_number not in skip_generic_invoice_no:
        _update_invoice_no_fields(doc, data.document_invoice_no)
    update_dated_in_document(doc, invoice_number=data.invoice_number)

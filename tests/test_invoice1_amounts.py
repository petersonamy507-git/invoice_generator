import sys
from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document

from backend.app.models import InvoiceTask, WordInvoiceData
from backend.app.services.word_field_updater import update_invoice_fields
from backend.app.services.word_invoice_generator import generate_word_invoice
from backend.app.services.word_templates import ensure_template_exists


def _invoice1_data() -> WordInvoiceData:
    return WordInvoiceData(
        person_name="John Smith",
        company_name="John Smith",
        email="t@t.com",
        total_amount=5000,
        invoice_number=1,
        tasks=[
            InvoiceTask(main_task="M", sub_task="Task Alpha", amount=2000),
            InvoiceTask(main_task="M", sub_task="Task Beta", amount=1500),
            InvoiceTask(main_task="M", sub_task="Task Gamma", amount=1500),
        ],
        bank="Test Bank",
        iban="PK00TEST",
        branch_code="1285",
        document_invoice_no="HH-008",
    )


def test_invoice1_line_items_and_total_from_sheet():
    doc = Document(str(ensure_template_exists(1)))
    update_invoice_fields(doc, _invoice1_data())

    table = doc.tables[1]
    assert table.rows[1].cells[0].text.strip() == "Task Alpha"
    assert table.rows[2].cells[0].text.strip() == "Task Beta"
    assert table.rows[3].cells[0].text.strip() == "Task Gamma"
    assert "$2,000" in table.rows[1].cells[1].text
    assert "$1,500" in table.rows[2].cells[1].text
    assert "$1,500" in table.rows[3].cells[1].text
    assert "$5,000" in table.rows[4].cells[1].text


def test_invoice1_generate_word_invoice():
    docx_bytes, _ = generate_word_invoice(_invoice1_data())
    doc = Document(BytesIO(docx_bytes))
    table = doc.tables[1]
    assert table.rows[1].cells[0].text.strip() == "Task Alpha"
    assert "$5,000" in table.rows[4].cells[1].text


if __name__ == "__main__":
    test_invoice1_line_items_and_total_from_sheet()
    test_invoice1_generate_word_invoice()
    print("Invoice 1 amount tests passed.")

import sys
from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document

from backend.app.models import InvoiceTask, WordInvoiceData
from backend.app.services.word_field_updater import update_invoice_fields
from backend.app.services.word_invoice_generator import generate_word_invoice
from backend.app.services.word_templates import ensure_template_exists


def _invoice4_data() -> WordInvoiceData:
    return WordInvoiceData(
        person_name="John Smith",
        company_name="John Smith",
        email="t@t.com",
        total_amount=5000,
        invoice_number=4,
        tasks=[
            InvoiceTask(main_task="M", sub_task="Task Alpha", amount=2000),
            InvoiceTask(main_task="M", sub_task="Task Beta", amount=1500),
            InvoiceTask(main_task="M", sub_task="Task Gamma", amount=1500),
        ],
        bank="Test Bank",
        iban="PK00TEST",
        branch_code="Lahore",
    )


def test_invoice4_preserves_header_layout_paragraphs():
    doc = Document(str(ensure_template_exists(4)))
    before_header = doc.tables[2].rows[0].cells[1].text
    update_invoice_fields(doc, _invoice4_data())
    assert doc.tables[2].rows[0].cells[1].text == before_header
    assert any("Invoice No." in (p.text or "") for p in doc.paragraphs)


def test_invoice4_updates_only_required_fields():
    docx_bytes, _ = generate_word_invoice(_invoice4_data())
    doc = Document(BytesIO(docx_bytes))

    payment = doc.tables[1].rows[0].cells[0].text
    assert "Test Bank" in payment
    assert "John Smith" in payment
    assert "PK00TEST" in payment
    assert "Lahore" in payment

    items = doc.tables[2]
    assert items.rows[1].cells[1].text.strip() == "Task Alpha"
    assert items.rows[2].cells[1].text.strip() == "Task Beta"
    assert items.rows[3].cells[1].text.strip() == "Task Gamma"
    assert "$2,000" in items.rows[1].cells[3].text
    assert "$5,000" in items.rows[4].cells[3].text
    assert "$5,000" in doc.tables[3].rows[0].cells[0].text


if __name__ == "__main__":
    test_invoice4_preserves_header_layout_paragraphs()
    test_invoice4_updates_only_required_fields()
    print("Invoice 4 layout tests passed.")

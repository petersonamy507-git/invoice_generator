import os
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document

from backend.app.services.html_to_docx import convert_html_to_docx_bytes
from backend.app.services.libreoffice import convert_with_libreoffice
from backend.app.services.validation import ValidationError
from backend.app.services.word_doc_converter import convert_doc_to_docx
from backend.app.services.word_to_pdf import convert_docx_batch_to_pdf


class LibreOfficeConversionTests(unittest.TestCase):
    @staticmethod
    def write_docx(path):
        document = Document()
        document.add_paragraph("Converted invoice")
        document.save(path)

    def test_runner_uses_headless_mode_and_isolated_profile(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            source.write_bytes(b"docx")
            output_dir = root / "output"

            with (
                patch.dict(os.environ, {"SOFFICE_PATH": "fake-soffice"}),
                patch(
                    "backend.app.services.libreoffice.subprocess.run",
                    return_value=SimpleNamespace(returncode=0, stderr="", stdout=""),
                ) as run,
            ):
                convert_with_libreoffice([source], output_dir, "pdf")

            command = run.call_args.args[0]
            self.assertEqual(command[0], "fake-soffice")
            self.assertIn("--headless", command)
            self.assertIn("--convert-to", command)
            self.assertEqual(command[command.index("--convert-to") + 1], "pdf")
            profile = next(arg for arg in command if arg.startswith("-env:UserInstallation="))
            self.assertIn("file:", profile)
            self.assertIn("libreoffice-profile-", profile)
            self.assertEqual(run.call_args.kwargs["timeout"], 300)

    def test_batch_conversion_preserves_pdf_names_and_bytes(self):
        def fake_convert(input_paths, output_dir, target_format):
            self.assertEqual(target_format, "pdf")
            for path in input_paths:
                (output_dir / f"{path.stem}.pdf").write_bytes(b"%PDF-test")

        with patch(
            "backend.app.services.word_to_pdf.convert_with_libreoffice",
            side_effect=fake_convert,
        ):
            result = convert_docx_batch_to_pdf(
                [(b"docx-one", "invoice-one.docx"), (b"docx-two", "invoice-two.docx")]
            )

        self.assertEqual(
            result,
            [
                (b"%PDF-test", "invoice-one.pdf"),
                (b"%PDF-test", "invoice-two.pdf"),
            ],
        )

    def test_html_conversion_returns_valid_docx_bytes(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            html_path = root / "invoice.html"
            html_path.write_text("<html><body>Invoice</body></html>", encoding="utf-8")

            def fake_convert(input_paths, output_dir, target_format):
                self.assertEqual(input_paths, [html_path.resolve()])
                self.assertEqual(target_format, "docx")
                self.write_docx(output_dir / "invoice.docx")

            with patch(
                "backend.app.services.html_to_docx.convert_with_libreoffice",
                side_effect=fake_convert,
            ):
                result = convert_html_to_docx_bytes(html_path)

        self.assertEqual(Document(BytesIO(result)).paragraphs[0].text, "Converted invoice")

    def test_legacy_doc_conversion_returns_valid_docx_path(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc_path = Path(tmp) / "legacy.doc"
            doc_path.write_bytes(b"legacy doc")

            def fake_convert(input_paths, output_dir, target_format):
                self.assertEqual(input_paths, [doc_path.resolve()])
                self.assertEqual(target_format, "docx")
                self.write_docx(output_dir / "legacy.docx")

            with patch(
                "backend.app.services.word_doc_converter.convert_with_libreoffice",
                side_effect=fake_convert,
            ):
                converted_path = convert_doc_to_docx(doc_path)
            try:
                self.assertEqual(
                    Document(str(converted_path)).paragraphs[0].text,
                    "Converted invoice",
                )
            finally:
                converted_path.unlink(missing_ok=True)
                converted_path.parent.rmdir()

    def test_nonzero_exit_reports_libreoffice_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            source.write_bytes(b"docx")
            with (
                patch.dict(os.environ, {"SOFFICE_PATH": "fake-soffice"}),
                patch(
                    "backend.app.services.libreoffice.subprocess.run",
                    return_value=SimpleNamespace(
                        returncode=1, stderr="conversion failed", stdout=""
                    ),
                ),
                self.assertRaisesRegex(ValidationError, "conversion failed"),
            ):
                convert_with_libreoffice([source], root / "output", "pdf")


if __name__ == "__main__":
    unittest.main()
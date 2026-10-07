import sys
from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document

from backend.app.models import InvoiceTask, WordInvoiceData
from backend.app.services.word_field_updater import update_invoice_fields
from backend.app.services.word_invoice_generator import generate_word_invoice
from backend.app.services.word_templates import ensure_template_exists


def _invoice3_data() -> WordInvoiceData:
    return WordInvoiceData(
        person_name="John Smith",
        company_name="John Smith",
        email="t@t.com",
        total_amount=5000,
        invoice_number=3,
        tasks=[
            InvoiceTask(main_task="M", sub_task="Task Alpha", amount=2000),
            InvoiceTask(main_task="M", sub_task="Task Beta", amount=1500),
            InvoiceTask(main_task="M", sub_task="Task Gamma", amount=1500),
        ],
        bank="Test Bank",
        iban="PK00TEST",
        branch_code="khushab",
    )


def test_invoice3_preserves_header_layout_paragraphs():
    doc = Document(str(ensure_template_exists(3)))
    before = {
        i: (doc.paragraphs[i].text, len(doc.paragraphs[i].runs))
        for i in (9, 10, 11)
        if i < len(doc.paragraphs)
    }
    before_header = doc.tables[1].rows[0].cells[0].text

    update_invoice_fields(doc, _invoice3_data())

    for i, (text, _runs) in before.items():
        assert doc.paragraphs[i].text == text
    assert doc.tables[1].rows[0].cells[0].text == before_header
    assert any("John Smith" in (p.text or "") for p in doc.paragraphs)


def test_invoice3_updates_only_required_fields():
    docx_bytes, _ = generate_word_invoice(_invoice3_data())
    doc = Document(BytesIO(docx_bytes))

    payment = doc.tables[3].rows[0].cells[0].text
    assert "Test Bank" in payment
    assert "PK00TEST" in payment
    assert "khushab" in payment

    items = doc.tables[1]
    assert items.rows[1].cells[0].text.strip() == "Task Alpha"
    assert items.rows[2].cells[0].text.strip() == "Task Beta"
    assert items.rows[3].cells[0].text.strip() == "Task Gamma"
    assert "$2,000" in items.rows[1].cells[1].text

    summary = doc.tables[2].rows[0].cells[0].text
    assert "$5,000" in summary
    assert "$0" in summary


if __name__ == "__main__":
    test_invoice3_preserves_header_layout_paragraphs()
    test_invoice3_updates_only_required_fields()
    print("Invoice 3 layout tests passed.")

import sys
from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document

from backend.app.models import InvoiceTask, WordInvoiceData
from backend.app.services.word_field_updater import update_invoice_fields
from backend.app.services.word_invoice_generator import generate_word_invoice
from backend.app.services.word_templates import ensure_template_exists


def _invoice2_data() -> WordInvoiceData:
    return WordInvoiceData(
        person_name="John Smith",
        company_name="John Smith",
        email="t@t.com",
        total_amount=5000,
        invoice_number=2,
        tasks=[
            InvoiceTask(main_task="M", sub_task="Task Alpha", amount=2000),
            InvoiceTask(main_task="M", sub_task="Task Beta", amount=1500),
            InvoiceTask(main_task="M", sub_task="Task Gamma", amount=1500),
        ],
        bank="Test Bank",
        iban="PK00TEST",
        branch_code="0213",
        document_invoice_no="SMDR-002",
    )


def test_invoice2_line_items_subtotal_and_total():
    doc = Document(str(ensure_template_exists(2)))
    update_invoice_fields(doc, _invoice2_data())

    table = doc.tables[1]
    assert table.rows[5].cells[0].text.strip() == "Task Alpha"
    assert table.rows[6].cells[0].text.strip() == "Task Beta"
    assert table.rows[7].cells[0].text.strip() == "Task Gamma"
    assert "$2,000" in table.rows[5].cells[2].text
    assert "$5,000" in table.rows[8].cells[2].text
    assert "$0" in table.rows[9].cells[2].text

    total_para = next(
        p.text for p in doc.paragraphs if p.text.strip().lower().startswith("total")
    )
    assert "$5,000" in total_para

    payment = doc.tables[2].rows[0].cells[0]
    payment_text = "\n".join(p.text for p in payment.paragraphs)
    assert "John Smith" in payment_text
    assert "PK00TEST" in payment_text
    assert "Test Bank" in payment_text
    assert "0213" in payment_text


def test_invoice2_generate_word_invoice():
    docx_bytes, _ = generate_word_invoice(_invoice2_data())
    doc = Document(BytesIO(docx_bytes))
    assert doc.tables[1].rows[5].cells[0].text.strip() == "Task Alpha"
    assert "$5,000" in doc.tables[1].rows[8].cells[2].text


if __name__ == "__main__":
    test_invoice2_line_items_subtotal_and_total()
    test_invoice2_generate_word_invoice()
    print("Invoice 2 amount tests passed.")

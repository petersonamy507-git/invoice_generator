import sys
from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import pandas as pd
from docx import Document

from backend.app.services.payment_sheet_processor import (
    build_word_invoices_from_payment_sheet,
)
from backend.app.services.word_field_updater import update_invoice_fields
from backend.app.services.word_templates import ensure_template_exists


def test_payment_sheet_updates_bank_fields():
    df = pd.DataFrame(
        [
            {
                "Name": "John Smith",
                "Bank": "Test Bank Ltd",
                "IBAN NO.": "PK00TEST123456789",
                "Branch Code": "9999",
                "Invoice": 4,
                "Invoice No.": "AA-001",
                "Amount": 5000,
            },
            {
                "Name": "Jane Doe",
                "Bank": "Another Bank",
                "IBAN NO.": "PK11OTHER987654321",
                "Branch Code": "BR-0213",
                "Invoice": 5,
                "Invoice No.": "DV-001",
                "Amount": 3000,
            },
        ]
    )
    buf = BytesIO()
    df.to_excel(buf, index=False)
    result = build_word_invoices_from_payment_sheet(buf.getvalue())
    assert len(result.invoices) == 2
    assert result.invoices[0].bank == "Test Bank Ltd"
    assert result.invoices[1].branch_code == "BR-0213"

    doc = Document(str(ensure_template_exists(4)))
    update_invoice_fields(doc, result.invoices[0])
    inv4_text = "\n".join(p.text for p in doc.paragraphs)
    inv4_text += "\n" + doc.tables[1].rows[0].cells[0].text
    assert "Test Bank Ltd" in inv4_text
    assert "John Smith" in inv4_text
    assert "PK00TEST123456789" in inv4_text
    assert "9999" in inv4_text

    doc5 = Document(str(ensure_template_exists(5)))
    update_invoice_fields(doc5, result.invoices[1])
    payment_block = doc5.paragraphs[21].text
    assert "Another Bank" in payment_block
    assert "Jane Doe" in payment_block
    assert "BR-0213" in doc5.paragraphs[22].text


def test_branch_code_on_all_invoice_templates():
    from backend.app.models import InvoiceTask, WordInvoiceData

    cases = [
        (1, "1285"),
        (2, "0213"),
        (3, "khushab"),
        (4, "Lahore"),
        (5, "Karachi"),
    ]
    for inv_num, branch in cases:
        data = WordInvoiceData(
            person_name="Test Person",
            company_name="Test Person",
            email="t@t.com",
            total_amount=10000,
            invoice_number=inv_num,
            tasks=[InvoiceTask(main_task="M", sub_task="A", amount=4000)] * 3,
            bank="My Bank",
            iban="PK00TEST",
            branch_code=branch,
        )
        doc = Document(str(ensure_template_exists(inv_num)))
        update_invoice_fields(doc, data)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text += "\n" + cell.text
        assert branch in full_text, f"Invoice {inv_num} missing branch code {branch!r}"


def test_branch_code_numeric_and_mixed_from_excel():
    from backend.app.services.excel_parser import _excel_text_value

    assert _excel_text_value(1285) == "1285"
    assert _excel_text_value(1285.0) == "1285"
    assert _excel_text_value("BR-0213") == "BR-0213"
    assert _excel_text_value("0213") == "0213"


if __name__ == "__main__":
    test_payment_sheet_updates_bank_fields()
    test_branch_code_on_all_invoice_templates()
    test_branch_code_numeric_and_mixed_from_excel()
    print("Payment sheet tests passed.")
